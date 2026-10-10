"""Сценарии отчётов по эффективности (ТЗ «Коэффициент эффективности», «Что получает начальник»; SPEC 7.8).

Бот целиком (bot.main.build_dispatcher) на фейковом Telegram API (tests/e2e/fakebot.py), AI выключен.
Задачи засеваются через сервисы — так, как их поставил бы начальник и сдал сотрудник; оценки
начальник подтверждает на экране проверки (или сервисом проверки), а отчёты смотрятся «глазами»
пользователей: тексты, кнопки, alert'ы, файлы.

Время бота заморожено фикстурой ``clock``: «сейчас» — четверг 08.10.2026 12:00 по Ташкенту, текущая
неделя — 05.10–11.10.2026, месяц — октябрь, квартал — IV. Подписи периодов и просрочки поэтому
не зависят от дня, в который запускаются тесты.
"""

from __future__ import annotations

import io
import sys
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timedelta

import pytest
from aiogram.methods import DeleteMessage, EditMessageText, SendDocument, SendMessage
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from bot.config import get_settings
from bot.db import models
from bot.db.models import ReviewDecision, Task, TaskSource, TaskStatus, User
from bot.services import kpi, periods, tasks, users
from bot.ui.callbacks import PeriodCB, UserCB
from bot.ui.texts import BTN_EXPORT, BTN_MY_KPI, BTN_PROPOSE, BTN_REVIEW, BTN_TEAM
from bot.utils.dates import to_utc

from .fakebot import MANAGER_TG_ID, BotHarness

pytestmark = pytest.mark.asyncio

MGR = MANAGER_TG_ID
IVANOV = 2001
SIDOROV = 2002
NOVIKOVA = 2003
NEWCOMER = 2009  # подал заявку, ещё не подтверждён

NO_RIGHTS = "⛔ Недостаточно прав для этого действия."

# «Сейчас» для бота — четверг замороженной недели, местное время Ташкента.
NOW_LOCAL = datetime(2026, 10, 8, 12, 0)
WEEK_LABEL = "Неделя 05.10–11.10.2026"
PREV_WEEK_LABEL = "Неделя 28.09–04.10.2026"


def local(day: int, month: int = 10, hour: int = 18, minute: int = 0, year: int = 2026) -> datetime:
    """Местное время Ташкента -> naive UTC (как в БД). По умолчанию — 18:00 октября 2026."""
    return to_utc(datetime(year, month, day, hour, minute))


# Пример из ТЗ: веса 30/20/20/30, итоговые оценки 100/110/90/105 -> 101,5 -> «102 %».
# (название, вес, оценка AI, итоговая оценка начальника, срок — день октября, 18:00)
TZ_TABLE = (
    ("Подготовка ТЗ", 30, 100, 100, 5),
    ("Анализ договоров", 20, 110, 110, 6),
    ("Отчёт", 20, 95, 90, 7),
    ("Работа с поставщиками", 30, 105, 105, 9),
)


# --- Часы и настройки ---------------------------------------------------------------------------


class Clock:
    """Часы бота: подменяют ``utcnow`` во всех модулях ``bot.*`` (naive UTC)."""

    def __init__(self, local_now: datetime) -> None:
        self.now = to_utc(local_now)

    def __call__(self) -> datetime:
        return self.now

    @contextmanager
    def at(self, moment: datetime) -> Iterator[None]:
        """Временно перевести часы (например, «сотрудник сдал результат во вторник»)."""
        saved, self.now = self.now, moment
        try:
            yield
        finally:
            self.now = saved


@pytest.fixture
def clock(app: BotHarness, monkeypatch: pytest.MonkeyPatch) -> Clock:
    """Заморозить время бота на четверге 08.10.2026 12:00 (Ташкент)."""
    frozen = Clock(NOW_LOCAL)
    original = models.utcnow
    for name, module in list(sys.modules.items()):
        if module is None or not (name == "bot" or name.startswith("bot.")):
            continue
        if getattr(module, "utcnow", None) is original:
            monkeypatch.setattr(module, "utcnow", frozen)
    # Ожидаемые числа считаются по настройкам по умолчанию — не зависеть от локального .env.
    monkeypatch.setenv("OVERDUE_COUNTS_AS_ZERO", "true")
    monkeypatch.setenv("MAX_SCORE", "150")
    get_settings.cache_clear()
    return frozen


# --- Засев данных -----------------------------------------------------------------------------


@dataclass
class Team:
    mgr: User
    ivanov: User
    sidorov: User
    novikova: User


async def seed_team(h: BotHarness) -> Team:
    """Начальник Петрова и три активных сотрудника (Новикова — без задач)."""
    return Team(
        mgr=await h.seed_user(MGR, "Петрова Анна Сергеевна", role="manager", position="Начальник управления"),
        ivanov=await h.seed_user(IVANOV, "Иванов Иван Иванович", position="Юрист"),
        sidorov=await h.seed_user(SIDOROV, "Сидоров Пётр Ильич", position="Экономист"),
        novikova=await h.seed_user(NOVIKOVA, "Новикова Елена Андреевна"),
    )


async def _user(session: AsyncSession, tg_id: int) -> User:
    user = await session.scalar(select(User).where(User.tg_id == tg_id))
    assert user is not None, tg_id
    return user


def _submit_moment(clock: Clock, deadline: datetime, late: bool) -> datetime:
    """Когда сотрудник сдал результат: в срок — за 3 часа до срока (но не позже, чем час назад);
    с опозданием — через 6 часов после срока. Проверка начальником — ещё через полчаса."""
    if late:
        moment = deadline + timedelta(hours=6)
    else:
        moment = min(deadline - timedelta(hours=3), clock.now - timedelta(hours=1))
    assert moment + timedelta(minutes=30) < clock.now, "засев не должен заглядывать в будущее"
    return moment


async def seed_active(
    h: BotHarness, clock: Clock, assignee: int, title: str, weight: int, deadline: datetime,
    *, by_employee: bool = False,
) -> int:
    """Задача в работе: поставлена начальником (или внесена сотрудником и подтверждена)
    за 3 дня до срока, но не позже, чем за 2 часа до «сейчас»."""
    async with h.db() as s:
        mgr, emp = await _user(s, MGR), await _user(s, assignee)
        with clock.at(min(deadline - timedelta(days=3), clock.now - timedelta(hours=2))):
            if by_employee:
                task = await tasks.propose_task(
                    s, employee=emp, title=title, expected_result=f"Результат: {title}", deadline=deadline
                )
                await tasks.approve_proposal(s, task.id, mgr, weight=weight)
            else:
                task = await tasks.create_task(
                    s, creator=mgr, assignee_id=emp.id, title=title,
                    expected_result=f"Результат: {title}", deadline=deadline, weight=weight,
                )
        await s.commit()
        return task.id


async def seed_submitted(
    h: BotHarness, clock: Clock, assignee: int, title: str, weight: int, deadline: datetime,
    *, ai: float, late: bool = False, by_employee: bool = False,
) -> tuple[int, int]:
    """Задача сдана сотрудником и получила предварительную оценку ``ai``. -> (task_id, sub_id)."""
    task_id = await seed_active(h, clock, assignee, title, weight, deadline, by_employee=by_employee)
    async with h.db() as s:
        emp = await _user(s, assignee)
        with clock.at(_submit_moment(clock, deadline, late)):
            sub = await tasks.submit_result(s, task_id, emp, fact_text=f"Выполнено: {title}")
            await tasks.record_evaluation(
                s, sub.id, score=ai, rationale="Расчёт по правилам (AI недоступен): план выполнен.", source="rules"
            )
        await s.commit()
        return task_id, sub.id


async def seed_done(
    h: BotHarness, clock: Clock, assignee: int, title: str, weight: int, score: float, deadline: datetime,
    *, ai: float | None = None, late: bool = False, by_employee: bool = False, rework: bool = False,
) -> int:
    """Задача выполнена: сдана, начальник подтвердил оценку AI (или поставил свою ``score``).

    rework=True: первую сдачу начальник вернул на доработку, через 10 минут сотрудник сдал снова.
    """
    ai = score if ai is None else ai
    task_id, sub_id = await seed_submitted(
        h, clock, assignee, title, weight, deadline, ai=ai, late=late, by_employee=by_employee
    )
    submitted = _submit_moment(clock, deadline, late)
    async with h.db() as s:
        mgr = await _user(s, MGR)
        if rework:
            with clock.at(submitted + timedelta(minutes=10)):
                await tasks.review_rework(s, sub_id, mgr, "Добавьте расчёты")
            with clock.at(submitted + timedelta(minutes=20)):
                sub = await tasks.submit_result(s, task_id, await _user(s, assignee), fact_text="Доработано")
                await tasks.record_evaluation(s, sub.id, score=ai, rationale="Повторная сдача.", source="rules")
                sub_id = sub.id
        with clock.at(submitted + timedelta(minutes=30)):
            if ai == score:
                await tasks.review_confirm(s, sub_id, mgr)
            else:
                await tasks.review_set_score(s, sub_id, mgr, score, "Оценка скорректирована")
        await s.commit()
    return task_id


async def seed_tz_table(h: BotHarness, clock: Clock, *, late_supplier: bool = False) -> None:
    """Иванову — четыре задачи из таблицы ТЗ, все проверены начальником.

    late_supplier=True: «Работа с поставщиками» внесена самим Ивановым и сдана с опозданием.
    """
    for title, weight, ai, final, day in TZ_TABLE:
        if late_supplier and title == "Работа с поставщиками":
            # Срок — среда 18:00, сдано в четверг ночью (на 6 ч позже), проверено до «сейчас».
            await seed_done(h, clock, IVANOV, title, weight, final, local(7), ai=ai, late=True, by_employee=True)
        else:
            await seed_done(h, clock, IVANOV, title, weight, final, local(day), ai=ai)


async def kpi_of(h: BotHarness, tg_id: int, kind: str = "week", offset: int = 0) -> kpi.KpiResult:
    """KPI сотрудника по сервису — для сверки с тем, что показал бот."""
    async with h.db() as s:
        user = await _user(s, tg_id)
        now = models.utcnow()
        return await kpi.kpi_for_user(s, user.id, periods.get_period(kind, offset, now), now)


def lines(text: str | None) -> list[str]:
    return [line.strip() for line in (text or "").splitlines()]


# --- 📊 Команда: пример из ТЗ ------------------------------------------------------------------


async def test_team_dashboard_shows_tz_example_after_review(app: BotHarness, clock: Clock) -> None:
    """Иванов сдал четыре задачи из таблицы ТЗ. Начальник Петрова в «📝 На проверке» три оценки
    подтверждает, а «Отчёт» (AI предложил 95 %) оценивает на 90 %. Затем открывает «📊 Команда»:
    у Иванова 102 % (101,5 по формуле), у Сидорова 80 %, у Новиковой задач нет, эффективность
    команды — взвешенно по всем оценённым задачам: (10150 + 50×80) / 150 ≈ 94 %."""
    h = app
    await seed_team(h)
    for title, weight, ai, _final, day in TZ_TABLE:
        await seed_submitted(h, clock, IVANOV, title, weight, local(day), ai=ai)
    await seed_done(h, clock, SIDOROV, "Бюджет на квартал", 50, 80, local(6))
    await h.send_command(MGR, "start")

    for title, _weight, ai, final, _day in TZ_TABLE:
        await h.press_menu(MGR, BTN_REVIEW)
        await h.press_button(MGR, title)
        if ai == final:
            log = await h.press_button(MGR, f"Подтвердить {ai} %")
            assert log.alert == "✅ Оценка подтверждена"
        else:
            await h.press_button(MGR, "Изменить оценку")
            await h.press_button(MGR, f"{final} %")
            log = await h.press_button(MGR, "Пропустить")
            assert log.alert == "✏️ Оценка сохранена"
        assert log.to(IVANOV).texts, "сотрудник получил итоговую оценку"

    done = await h.scalars(select(Task).where(Task.assignee_id == (await h.get_user(IVANOV)).id))
    assert {t.title: (t.status, t.final_score) for t in done} == {
        title: (TaskStatus.DONE, float(final)) for title, _w, _ai, final, _d in TZ_TABLE
    }
    decisions = {t.title: t.last_submission.decision for t in done}
    assert decisions["Отчёт"] == ReviewDecision.CHANGED
    assert decisions["Подготовка ТЗ"] == ReviewDecision.APPROVED
    assert (await kpi_of(h, IVANOV)).kpi == 101.5

    log = await h.press_menu(MGR, BTN_TEAM)
    assert log.of(SendMessage), "дашборд пришёл новым сообщением"
    text = h.last_text(MGR)
    rows = lines(text)
    assert f"📊 Команда · {WEEK_LABEL}" in rows
    assert "Эффективность команды: 94 %" in rows
    assert "Итого: ✅ 5 · ⏰ 0 · 🔄 0 · 📝 0" in rows
    ivanov_at = rows.index("1. Иванов И. И. — 102 %")
    assert rows[ivanov_at + 1].endswith("✅ 4 · ⏰ 0 · 🔄 0 · 📝 0")
    assert "2. Сидоров П. И. — 80 %" in rows
    novikova_at = rows.index("3. Новикова Е. А. — нет данных")
    assert rows[novikova_at + 1] == "задач нет"
    assert h.buttons(MGR) == [
        "• Неделя", "Месяц", "Квартал", "Год",
        "◀ Раньше",
        "👤 Иванов И. И. — 102 %",
        "👤 Сидоров П. И. — 80 %",
        "👤 Новикова Е. А.",
    ]


async def test_team_dashboard_without_employees(app: BotHarness, clock: Clock) -> None:
    """В команде ещё никого не подтвердили: начальник видит «Активных сотрудников пока нет»
    и «нет данных», а не пустой экран или ошибку."""
    h = app
    await h.seed_user(MGR, "Петрова Анна Сергеевна", role="manager")
    await h.seed_user(NEWCOMER, "Новичков Олег Петрович", status="pending")
    await h.send_command(MGR, "team")
    rows = lines(h.last_text(MGR))
    assert "Эффективность команды: нет данных" in rows
    assert "Активных сотрудников пока нет." in rows
    assert not any("Новичков" in button for button in h.buttons(MGR))


async def test_large_team_fits_one_message(app: BotHarness, clock: Clock) -> None:
    """В управлении 70 сотрудников с длинными ФИО. Дашборд помещается в одно сообщение Telegram:
    лишние строки сворачиваются в «… и ещё N сотрудников», кнопок — не больше лимита Telegram (100)."""
    h = app
    await h.seed_user(MGR, "Петрова Анна Сергеевна", role="manager")
    for n in range(70):
        await h.seed_user(3000 + n, f"Константинопольская-Преображенская{n:02d} Александра Владимировна")
    await seed_done(h, clock, 3000, "Подготовка ТЗ", 30, 100, local(5))

    await h.send_command(MGR, "team")  # FakeSession сама отклонила бы текст длиннее 4096 символов
    rows = lines(h.last_text(MGR))
    assert "1. Константинопольская-Преображенская00 А. В. — 100 %" in rows
    assert any(row.startswith("… и ещё") and row.endswith("сотрудников") for row in rows)
    assert rows[-1] == "✅ выполнено · ⏰ просрочено · 🔄 в работе · 📝 на проверке"
    assert len(h.buttons(MGR)) <= 100


# --- Периоды: неделя / месяц / квартал / год, листание ------------------------------------------


async def test_period_switching_and_navigation(app: BotHarness, clock: Clock) -> None:
    """Начальник листает дашборд команды. У Иванова три оценённые задачи: на этой неделе (100 %),
    на прошлой — 1 октября (80 %) и 23 сентября (60 %). Неделя/месяц/квартал/год считаются той же
    формулой по задачам периода; ◀ листает назад, ▶ есть только у прошлых периодов; всё — правкой
    одного и того же сообщения."""
    h = app
    await seed_team(h)
    await seed_done(h, clock, IVANOV, "Текущая неделя", 20, 100, local(6))
    await seed_done(h, clock, IVANOV, "Прошлая неделя", 20, 80, local(1))
    await seed_done(h, clock, IVANOV, "Сентябрьская задача", 20, 60, local(23, month=9))

    await h.send_command(MGR, "team")
    dashboard = h.last_message(MGR)
    assert dashboard is not None

    def head() -> tuple[str, str]:
        rows = lines(h.last_text(MGR))
        ivanov = next(row for row in rows if row.startswith("1. Иванов И. И."))
        return rows[0], ivanov

    async def press(text: str) -> None:
        log = await h.press_button(MGR, text, dashboard.message_id)
        assert log.answers, "callback.answer() вызван"
        assert log.of(EditMessageText) and not log.of(SendMessage), "период переключается правкой сообщения"

    assert head() == (f"📊 Команда · {WEEK_LABEL}", "1. Иванов И. И. — 100 %")
    assert not h.has_button(MGR, "Позже", dashboard.message_id), "в будущее листать нельзя"

    await press("◀ Раньше")
    assert head() == (f"📊 Команда · {PREV_WEEK_LABEL}", "1. Иванов И. И. — 80 %")
    assert "• Неделя" in h.buttons(MGR) and "Позже ▶" in h.buttons(MGR)

    await press("◀ Раньше")
    assert head() == ("📊 Команда · Неделя 21.09–27.09.2026", "1. Иванов И. И. — 60 %")
    await press("Позже ▶")
    assert head()[0] == f"📊 Команда · {PREV_WEEK_LABEL}"
    await press("Позже ▶")
    assert head()[0] == f"📊 Команда · {WEEK_LABEL}"
    assert "Позже ▶" not in h.buttons(MGR)

    expected = [
        ("Месяц", "Октябрь 2026", "90 %"),
        ("◀ Раньше", "Сентябрь 2026", "60 %"),
        ("Квартал", "IV квартал 2026", "90 %"),  # смена вида периода сбрасывает листание
        ("◀ Раньше", "III квартал 2026", "60 %"),
        ("Год", "2026 год", "80 %"),
    ]
    for button, label, value in expected:
        await press(button)
        assert head() == (f"📊 Команда · {label}", f"1. Иванов И. И. — {value}"), button
    assert "• Год" in h.buttons(MGR)

    await press("◀ Раньше")
    rows = lines(h.last_text(MGR))
    assert rows[0] == "📊 Команда · 2025 год"
    assert "Эффективность команды: нет данных" in rows
    assert rows[rows.index("1. Иванов И. И. — нет данных") + 1] == "задач нет"
    assert len(h.messages(MGR)) == 1, "все переключения — правки одного сообщения дашборда"


async def test_forged_period_callbacks_are_normalized(app: BotHarness, clock: Clock) -> None:
    """Подделанные кнопки периода: «будущая» неделя и неизвестный вид периода показывают текущую
    неделю, а не падают и не открывают будущее."""
    h = app
    await seed_team(h)
    await seed_done(h, clock, IVANOV, "Подготовка ТЗ", 30, 100, local(5))
    await h.send_command(MGR, "team")

    log = await h.press(MGR, PeriodCB(scope="team", kind="week", offset=3))
    assert log.answers
    assert lines(h.last_text(MGR))[0] == f"📊 Команда · {WEEK_LABEL}"
    assert "Позже ▶" not in h.buttons(MGR)

    await h.press(MGR, PeriodCB(scope="team", kind="month", offset=-1))
    await h.press(MGR, PeriodCB(scope="team", kind="decade", offset=0))
    assert lines(h.last_text(MGR))[0] == f"📊 Команда · {WEEK_LABEL}"
    assert "• Неделя" in h.buttons(MGR)

    # Очень далёкое прошлое тоже не роняет бота.
    log = await h.press(MGR, PeriodCB(scope="team", kind="year", offset=-100000))
    assert log.answers
    assert "Эффективность команды: нет данных" in lines(h.last_text(MGR))


async def test_blocked_manager_cannot_use_old_dashboard_buttons(app: BotHarness, clock: Clock) -> None:
    """Второго начальника Кузнецова заблокировали, а у него остался открытый дашборд.
    Кнопки периода, карточки и экспорта отвечают «Недостаточно прав» и ничего не показывают."""
    h = app
    team = await seed_team(h)
    await h.seed_user(1002, "Кузнецов Игорь Олегович", role="manager")
    await seed_done(h, clock, IVANOV, "Подготовка ТЗ", 30, 100, local(5))
    await h.send_command(1002, "team")
    dashboard = h.last_text(1002)
    assert "1. Иванов И. И. — 100 %" in lines(dashboard)

    async with h.db() as s:
        await users.block_user(s, (await _user(s, 1002)).id, await _user(s, MGR))
        await s.commit()
    for data in (
        h.find_button(1002, "◀ Раньше"),
        h.find_button(1002, "👤 Иванов"),
        PeriodCB(scope="export", kind="week"),
        UserCB(action="history", user_id=team.ivanov.id),
    ):
        log = await h.press(1002, data)
        assert log.alert == NO_RIGHTS, data
        assert not log.texts and not log.documents, data
    assert h.last_text(1002) == dashboard


async def test_week_boundaries_follow_local_time(app: BotHarness, clock: Clock) -> None:
    """Неделя считается по ташкентскому времени (с понедельника 00:00), а не по UTC. Срок
    «понедельник 05.10, 00:30» (в UTC это ещё воскресенье) — текущая неделя; «воскресенье 04.10,
    23:30» — прошлая; «воскресенье 11.10, 23:00» — ещё текущая."""
    h = app
    await seed_team(h)
    await seed_done(h, clock, IVANOV, "Срок в понедельник ночью", 10, 100, local(5, hour=0, minute=30))
    await seed_done(h, clock, IVANOV, "Срок в воскресенье ночью", 10, 50, local(4, hour=23, minute=30))
    await seed_active(h, clock, IVANOV, "Срок в конце недели", 10, local(11, hour=23))

    await h.send_command(MGR, "team")
    rows = lines(h.last_text(MGR))
    ivanov_at = rows.index("1. Иванов И. И. — 100 %")
    assert rows[ivanov_at + 1].endswith("✅ 1 · ⏰ 0 · 🔄 1 · 📝 0")

    await h.press_button(MGR, "◀ Раньше")
    rows = lines(h.last_text(MGR))
    assert rows[0] == f"📊 Команда · {PREV_WEEK_LABEL}"
    ivanov_at = rows.index("1. Иванов И. И. — 50 %")
    assert rows[ivanov_at + 1].endswith("✅ 1 · ⏰ 0 · 🔄 0 · 📝 0")


# --- 👤 Карточка сотрудника ----------------------------------------------------------------------


async def test_employee_card_shows_all_indicators(app: BotHarness, clock: Clock) -> None:
    """Из дашборда начальник открывает карточку Иванова: эффективность за неделю и месяц,
    «Выполнено задач», «Просрочено», «Выполнение в срок», «Перевыполнено», «Внесено самостоятельно»
    и что вошло в расчёт. «Работу с поставщиками» Иванов внёс сам и сдал с опозданием."""
    h = app
    await seed_team(h)
    await seed_tz_table(h, clock, late_supplier=True)
    await h.send_command(MGR, "team")

    log = await h.press_button(MGR, "👤 Иванов И. И.")
    assert log.answers and log.of(EditMessageText)
    rows = lines(h.last_text(MGR))
    assert rows[:3] == ["👤 Иванов Иван Иванович — 102 %", "💼 Юрист", f"📅 {WEEK_LABEL}"]
    for expected in (
        "✅ Выполнено задач: 4 из 4",
        "⏰ Просрочено: 1",
        "🎯 Выполнение в срок: 75 %",
        "🚀 Перевыполнено: 2",
        "✋ Внесено самостоятельно: 1",
        "🔄 В работе: 0 · 📝 На проверке: 0",
        "• Подготовка ТЗ — вес 30 % × 100 %",
        "• Анализ договоров — вес 20 % × 110 %",
        "• Отчёт — вес 20 % × 90 %",
        "• Работа с поставщиками — вес 30 % × 105 %",
        "📊 Неделя: 102 % · Месяц: 102 %",
    ):
        assert expected in rows, expected
    assert h.buttons(MGR) == [
        "• Неделя", "Месяц", "Квартал", "Год", "◀ Раньше", "📜 История оценок", "📋 Задачи", "◀ К команде",
    ]
    async with h.db() as s:
        supplier = await s.scalar(select(Task).where(Task.title == "Работа с поставщиками"))
        assert supplier.source == TaskSource.EMPLOYEE and supplier.last_submission.is_late

    # Выбранный период — месяц, затем прошлый месяц: неделя и месяц внизу карточки остаются текущими.
    await h.press_button(MGR, "Месяц")
    rows = lines(h.last_text(MGR))
    assert rows[0] == "👤 Иванов Иван Иванович — 102 %" and "📅 Октябрь 2026" in rows
    await h.press_button(MGR, "◀ Раньше")
    rows = lines(h.last_text(MGR))
    assert rows[0] == "👤 Иванов Иван Иванович — нет данных"
    assert "📅 Сентябрь 2026" in rows and "Задач в этом периоде нет." in rows
    assert "📊 Неделя: 102 % · Месяц: 102 %" in rows
    assert "Позже ▶" in h.buttons(MGR)

    # «◀ К команде» возвращает дашборд за тот же период.
    await h.press_button(MGR, "К команде")
    assert lines(h.last_text(MGR))[0] == "📊 Команда · Сентябрь 2026"

    # «📋 Задачи» из карточки — список задач именно этого сотрудника.
    await h.press(MGR, UserCB(action="card", user_id=(await h.get_user(IVANOV)).id))
    await h.press_button(MGR, "📋 Задачи")
    text = h.last_text(MGR) or ""
    assert "Подготовка ТЗ" in text and "Бюджет" not in text


async def test_employee_card_counts_overdue_as_zero_and_skips_review(app: BotHarness, clock: Clock) -> None:
    """У Иванова на неделе: выполненная задача (вес 30, 100 %), просроченная несданная (вес 20),
    задача на проверке (вес 20, AI предлагает 120 %) и задача в работе со сроком в пятницу (вес 30).
    В KPI входят только выполненная и просроченная (как 0 %): (30×100 + 20×0) / 50 = 60 %.
    Когда Иванов сдаёт просроченную задачу, она уходит «на проверку» и больше не тянет KPI вниз,
    но остаётся в «Просрочено»; после подтверждения 120 % KPI = (3000 + 2400) / 50 = 108 %."""
    h = app
    await seed_team(h)
    await seed_done(h, clock, IVANOV, "Подготовка ТЗ", 30, 100, local(5))
    overdue_id = await seed_active(h, clock, IVANOV, "Анализ договоров", 20, local(6))
    _, review_sub = await seed_submitted(h, clock, IVANOV, "Отчёт", 20, local(7), ai=120)
    await seed_active(h, clock, IVANOV, "Работа с поставщиками", 30, local(9))
    ivanov_id = (await h.get_user(IVANOV)).id

    await h.send_command(MGR, "team")
    rows = lines(h.last_text(MGR))
    ivanov_at = rows.index("1. Иванов И. И. — 60 %")
    assert rows[ivanov_at + 1].endswith("✅ 1 · ⏰ 1 · 🔄 1 · 📝 1")

    await h.press(MGR, UserCB(action="card", user_id=ivanov_id))
    rows = lines(h.last_text(MGR))
    assert rows[0] == "👤 Иванов Иван Иванович — 60 %"
    for expected in (
        "✅ Выполнено задач: 1 из 4",
        "⏰ Просрочено: 1",
        "🎯 Выполнение в срок: 100 %",
        "🔄 В работе: 1 · 📝 На проверке: 1",
        "• Подготовка ТЗ — вес 30 % × 100 %",
        "• Анализ договоров — вес 20 % × 0 % ⏰ просрочена",
    ):
        assert expected in rows, expected
    assert not any(row.startswith(("• Отчёт", "• Работа с поставщиками")) for row in rows)

    # Иванов сдаёт просроченную задачу (на сутки позже срока) -> на проверке, в KPI не входит.
    async with h.db() as s:
        sub = await tasks.submit_result(s, overdue_id, await _user(s, IVANOV), fact_text="Проверено 100 договоров")
        await tasks.record_evaluation(s, sub.id, score=96, rationale="С опозданием.", source="rules")
        await s.commit()
    await h.press(MGR, UserCB(action="card", user_id=ivanov_id))
    rows = lines(h.last_text(MGR))
    assert rows[0] == "👤 Иванов Иван Иванович — 100 %"
    assert "⏰ Просрочено: 1" in rows and "🔄 В работе: 1 · 📝 На проверке: 2" in rows

    # Начальник подтверждает 120 % за «Отчёт».
    async with h.db() as s:
        await tasks.review_confirm(s, review_sub, await _user(s, MGR))
        await s.commit()
    await h.press(MGR, PeriodCB(scope="emp", kind="week", offset=0, user_id=ivanov_id))
    rows = lines(h.last_text(MGR))
    assert rows[0] == "👤 Иванов Иван Иванович — 108 %"
    assert "🚀 Перевыполнено: 1" in rows


async def test_overdue_not_counted_when_setting_disabled(
    app: BotHarness, clock: Clock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Если в настройках OVERDUE_COUNTS_AS_ZERO=false, просроченная несданная задача в KPI не входит,
    но в «Просрочено» по-прежнему видна."""
    h = app
    await seed_team(h)
    await seed_done(h, clock, IVANOV, "Подготовка ТЗ", 30, 100, local(5))
    await seed_active(h, clock, IVANOV, "Анализ договоров", 20, local(6))
    monkeypatch.setenv("OVERDUE_COUNTS_AS_ZERO", "false")
    get_settings.cache_clear()

    await h.send_command(MGR, "team")
    assert "1. Иванов И. И. — 100 %" in lines(h.last_text(MGR))
    await h.press_button(MGR, "👤 Иванов")
    rows = lines(h.last_text(MGR))
    assert rows[0] == "👤 Иванов Иван Иванович — 100 %"
    assert "⏰ Просрочено: 1" in rows


async def test_card_of_unknown_user_and_inaccessible_message(app: BotHarness, clock: Clock) -> None:
    """Начальник нажимает кнопку карточки удалённого из базы сотрудника — alert «Сотрудник не найден.»;
    кнопка периода на старом (недоступном) сообщении — карточка приходит новым сообщением."""
    h = app
    team = await seed_team(h)
    await h.send_command(MGR, "team")

    log = await h.press(MGR, UserCB(action="card", user_id=999))
    assert log.alert == "Сотрудник не найден."
    log = await h.press(MGR, UserCB(action="history", user_id=999))
    assert log.alert == "Сотрудник не найден."

    dashboard = h.last_message(MGR)
    assert dashboard is not None
    await h.bot.delete_message(MGR, dashboard.message_id)
    log = await h.press(MGR, PeriodCB(scope="emp", kind="month", offset=0, user_id=team.sidorov.id))
    assert log.answers and log.of(SendMessage)
    assert lines(h.last_text(MGR))[0] == "👤 Сидоров Пётр Ильич — нет данных"


# --- 📜 История оценок ---------------------------------------------------------------------------


async def test_evaluation_history_pagination(app: BotHarness, clock: Clock) -> None:
    """У Иванова 12 оценённых задач (по одной в день с 26.09 по 07.10) и одна на проверке.
    История показывает по 10, свежие сверху, с оценкой AI → итоговой, решением начальника,
    опозданием и числом доработок; «Вперёд ▶» / «◀ Назад» листают, «📊 К карточке» возвращает карточку."""
    h = app
    await seed_team(h)
    days = [(26 + i, 9) if i < 5 else (i - 4, 10) for i in range(12)]  # 26.09 … 30.09, 01.10 … 07.10
    for n, (day, month) in enumerate(days, start=1):
        changed = n == 12
        await seed_done(
            h, clock, IVANOV, f"Задача {n:02d}", 10, 90 if changed else 100, local(day, month),
            ai=95 if changed else None, late=n == 11, rework=n == 10,
        )
    await seed_submitted(h, clock, IVANOV, "Ещё на проверке", 10, local(8), ai=100)
    ivanov_id = (await h.get_user(IVANOV)).id

    await h.send_command(MGR, "team")
    await h.press_button(MGR, "👤 Иванов")
    log = await h.press_button(MGR, "📜 История оценок")
    assert log.answers and log.of(EditMessageText)
    text = h.last_text(MGR) or ""
    rows = lines(text)
    assert rows[0] == "📜 История оценок · Иванов И. И."
    assert "Оценено задач: 12 · стр. 1 из 2" in rows
    titles = [row.split(" ", 1)[1] for row in rows if row.startswith("#")]
    assert titles == [f"Задача {n:02d}" for n in range(12, 2, -1)], "свежие сверху, по 10 на странице"
    assert "Ещё на проверке" not in text
    assert "🤖 95 % → 🏁 90 % ✏️ изменена" in rows
    assert "🤖 100 % → 🏁 100 % ✅ подтверждена" in rows
    assert any("⚠️ с опозданием" in row for row in rows)
    assert any("↩️ доработок: 1" in row for row in rows)
    assert h.buttons(MGR) == ["Вперёд ▶", "📊 К карточке"]

    await h.press_button(MGR, "Вперёд ▶")
    rows = lines(h.last_text(MGR))
    assert "Оценено задач: 12 · стр. 2 из 2" in rows
    assert [row.split(" ", 1)[1] for row in rows if row.startswith("#")] == ["Задача 02", "Задача 01"]
    assert h.buttons(MGR) == ["◀ Назад", "📊 К карточке"]

    await h.press_button(MGR, "◀ Назад")
    assert "Оценено задач: 12 · стр. 1 из 2" in lines(h.last_text(MGR))

    # Подделанная страница за пределами истории -> последняя страница.
    await h.press(MGR, UserCB(action="history", user_id=ivanov_id, page=99))
    assert "Оценено задач: 12 · стр. 2 из 2" in lines(h.last_text(MGR))

    await h.press_button(MGR, "📊 К карточке")
    assert lines(h.last_text(MGR))[0].startswith("👤 Иванов Иван Иванович — ")
    assert "◀ К команде" in h.buttons(MGR)


async def test_evaluation_history_empty(app: BotHarness, clock: Clock) -> None:
    """У Новиковой нет оценённых задач — история честно говорит об этом, листать нечего."""
    h = app
    team = await seed_team(h)
    await seed_active(h, clock, NOVIKOVA, "Подготовить договор", 20, local(9))
    await h.send_command(MGR, "team")
    await h.press(MGR, UserCB(action="history", user_id=team.novikova.id))
    rows = lines(h.last_text(MGR))
    assert rows[0] == "📜 История оценок · Новикова Е. А."
    assert "Оценённых задач пока нет." in rows
    assert h.buttons(MGR) == ["📊 К карточке"]


# --- 📈 Моя эффективность (сотрудник) ------------------------------------------------------------


async def test_my_kpi_shows_only_own_data(app: BotHarness, clock: Clock) -> None:
    """Иванов открывает «📈 Моя эффективность»: своя карточка с 102 % из примера ТЗ, переключение
    периодов, история оценок — и ни слова о коллегах."""
    h = app
    team = await seed_team(h)
    await seed_tz_table(h, clock)
    await seed_done(h, clock, SIDOROV, "Бюджет на квартал", 50, 80, local(6))
    await h.send_command(IVANOV, "start")

    log = await h.press_menu(IVANOV, BTN_MY_KPI)
    assert log.of(SendMessage)
    text = h.last_text(IVANOV) or ""
    rows = lines(text)
    assert rows[:3] == ["👤 Иванов Иван Иванович — 102 %", "💼 Юрист", f"📅 {WEEK_LABEL}"]
    assert "✅ Выполнено задач: 4 из 4" in rows and "📊 Неделя: 102 % · Месяц: 102 %" in rows
    assert "Сидоров" not in text and "Бюджет" not in text
    assert h.buttons(IVANOV) == [
        "• Неделя", "Месяц", "Квартал", "Год", "◀ Раньше", "📜 История оценок", "📋 Мои задачи",
    ]

    await h.press_button(IVANOV, "Квартал")
    rows = lines(h.last_text(IVANOV))
    assert rows[0] == "👤 Иванов Иван Иванович — 102 %" and "📅 IV квартал 2026" in rows
    await h.press_button(IVANOV, "◀ Раньше")
    assert "📅 III квартал 2026" in lines(h.last_text(IVANOV))

    # Подделанный PeriodCB("me") с чужим user_id всё равно показывает свою карточку.
    log = await h.press(IVANOV, PeriodCB(scope="me", kind="week", offset=0, user_id=team.sidorov.id))
    assert log.answers and not log.alert
    assert lines(h.last_text(IVANOV))[0] == "👤 Иванов Иван Иванович — 102 %"

    # Своя история оценок и возврат к своей карточке (без кнопок начальника).
    await h.press_button(IVANOV, "📜 История оценок")
    assert lines(h.last_text(IVANOV))[0] == "📜 История оценок · Иванов И. И."
    await h.press_button(IVANOV, "📊 К карточке")
    assert lines(h.last_text(IVANOV))[0] == "👤 Иванов Иван Иванович — 102 %"
    assert "◀ К команде" not in h.buttons(IVANOV) and "📋 Мои задачи" in h.buttons(IVANOV)

    # «📋 Мои задачи» из карточки — свои задачи, без чужих.
    log = await h.press_button(IVANOV, "📋 Мои задачи")
    assert log.answers and "Бюджет на квартал" not in (h.last_text(IVANOV) or "")

    # Сидоров видит свои 80 %.
    await h.send_command(SIDOROV, "kpi")
    assert lines(h.last_text(SIDOROV))[0] == "👤 Сидоров Пётр Ильич — 80 %"


async def test_employee_cannot_open_colleagues_or_team(app: BotHarness, clock: Clock) -> None:
    """Иванов подделывает кнопки: карточка, история и периоды Сидорова, дашборд команды, экспорт.
    Каждый раз — alert «Недостаточно прав», сообщение не меняется, данных Сидорова Иван не видит.
    Кнопка «📊 Команда» и /team, /export для сотрудника тоже не работают."""
    h = app
    team = await seed_team(h)
    await seed_done(h, clock, SIDOROV, "Бюджет на квартал", 50, 80, local(6))
    await h.send_command(IVANOV, "kpi")
    before = h.last_text(IVANOV)

    forged = [
        UserCB(action="card", user_id=team.sidorov.id),
        UserCB(action="history", user_id=team.sidorov.id),
        PeriodCB(scope="emp", kind="week", offset=0, user_id=team.sidorov.id),
        PeriodCB(scope="team", kind="week", offset=0),
        PeriodCB(scope="export", kind="week", offset=0),
    ]
    for data in forged:
        log = await h.press(IVANOV, data)
        assert log.alert == NO_RIGHTS, data
        assert not log.texts and not log.documents, data
    assert h.last_text(IVANOV) == before

    for text in (BTN_TEAM, "/team", BTN_EXPORT, "/export"):
        log = await h.send_text(IVANOV, text)
        assert "Команда ·" not in log.text and "Экспорт в Excel" not in log.text, text
        assert "Сидоров" not in log.text and not log.documents, text

    # Неподтверждённый пользователь не видит даже «свою» эффективность.
    await h.seed_user(NEWCOMER, "Новичков Олег Петрович", status="pending")
    log = await h.press(NEWCOMER, PeriodCB(scope="me", kind="week"))
    assert log.alert == NO_RIGHTS and not log.texts
    log = await h.send_text(NEWCOMER, BTN_MY_KPI)
    assert "Новичков Олег Петрович —" not in log.text


# --- 📤 Экспорт ---------------------------------------------------------------------------------


def _sheet_rows(data: bytes, title: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(data))
    return [tuple(row) for row in wb[title].iter_rows(values_only=True)]


def _records(rows: list[tuple]) -> list[dict]:
    """Строки листа (без шапки) как словари «колонка -> значение»; пустые строки пропускаются."""
    header = rows[0]
    return [dict(zip(header, row, strict=True)) for row in rows[1:] if any(cell is not None for cell in row)]


async def test_export_sends_xlsx_with_three_sheets(app: BotHarness, clock: Clock) -> None:
    """Начальник выбирает «📤 Экспорт» → «Эта неделя»: приходит Excel-файл kpi_week_20261005.xlsx,
    он открывается, в нём листы «Сводка», «Задачи», «Журнал»; KPI Иванова — 101,5 (числом),
    команды — 94,3; сообщение «⏳ Готовлю отчёт…» после отправки исчезает."""
    h = app
    await seed_team(h)
    await seed_tz_table(h, clock)
    await seed_done(h, clock, SIDOROV, "Бюджет на квартал", 50, 80, local(6))
    await h.send_command(MGR, "start")

    await h.press_menu(MGR, BTN_EXPORT)
    assert "Экспорт в Excel" in (h.last_text(MGR) or "")
    assert h.buttons(MGR) == ["Эта неделя", "Прошлая неделя", "Этот месяц", "Прошлый месяц", "Квартал", "Год"]

    log = await h.press_button(MGR, "Эта неделя")
    assert log.answers
    assert any(text.startswith(f"⏳ Готовлю отчёт: {WEEK_LABEL}") for text in log.texts)
    progress_id = next(m.message_id for m in log.of(DeleteMessage))
    assert all(msg.message_id != progress_id for msg in h.messages(MGR)), "сообщение о подготовке удалено"
    [doc] = log.documents
    assert doc.file_name == "kpi_week_20261005.xlsx"
    assert doc.caption == f"📊 Отчёт: {WEEK_LABEL}"
    assert doc.content

    wb = load_workbook(io.BytesIO(doc.content))
    assert wb.sheetnames == ["Сводка", "Задачи", "Журнал"]
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2" and ws["A1"].font.bold, ws.title
    summary = _sheet_rows(doc.content, "Сводка")
    header = summary[0]
    assert header == (
        "Сотрудник", "Должность", "KPI %", "Задач", "Выполнено", "В срок %", "Просрочено",
        "На проверке", "В работе", "Перевыполнено", "Внесено самостоятельно",
    )
    by_name = {row["Сотрудник"]: row for row in _records(summary) if row["Сотрудник"]}
    ivanov = by_name["Иванов Иван Иванович"]
    assert ivanov["KPI %"] == 101.5
    assert (ivanov["Задач"], ivanov["Выполнено"], ivanov["Перевыполнено"]) == (4, 4, 2)
    assert by_name["Сидоров Пётр Ильич"]["KPI %"] == 80
    assert by_name["Новикова Елена Андреевна"]["KPI %"] is None
    assert by_name["Итого по команде"]["KPI %"] == 94.3
    assert any(str(row[0]).startswith(f"Период: {WEEK_LABEL}") for row in summary if row and row[0])

    task_rows = _sheet_rows(doc.content, "Задачи")
    task_header = task_rows[0]
    assert task_header == (
        "№", "Сотрудник", "Задача", "Ожидаемый результат", "План", "Факт", "Срок", "Сдано", "Просрочка дн.",
        "Вес %", "Приоритет", "Статус", "Оценка AI", "Итоговая оценка", "Решение", "Комментарий",
    )
    titles = {row["Задача"]: row for row in _records(task_rows)}
    assert set(titles) == {t for t, *_ in TZ_TABLE} | {"Бюджет на квартал"}
    assert titles["Отчёт"]["Итоговая оценка"] == 90
    assert titles["Отчёт"]["Оценка AI"] == 95

    journal = _sheet_rows(doc.content, "Журнал")
    assert journal[0] == ("Дата", "Задача", "Кто", "Событие", "Детали")
    events = {row[3] for row in journal[1:]}
    assert {"Задача поставлена", "Сдан результат", "Оценка подтверждена", "Оценка изменена"} <= events

    # Прошлая неделя пустая — файл всё равно приходит, задач в нём нет.
    log = await h.press_button(MGR, "Прошлая неделя")
    [doc] = log.documents
    assert doc.file_name == "kpi_week_20260928.xlsx"
    assert len(_sheet_rows(doc.content, "Задачи")) == 1, "только шапка"
    prev = {row[0]: row[2] for row in _sheet_rows(doc.content, "Сводка")[1:] if row and row[0]}
    assert prev["Иванов Иван Иванович"] is None

    log = await h.press_button(MGR, "Год")
    [doc] = log.documents
    assert doc.file_name == "kpi_year_20260101.xlsx" and doc.caption == "📊 Отчёт: 2026 год"
    assert len(h.documents_sent(MGR)) == 3


async def test_export_survives_report_failure(
    app: BotHarness, clock: Clock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Если собрать отчёт не удалось (ошибка в выгрузке), начальник видит понятное сообщение
    вместо «⏳ Готовлю отчёт…», а бот продолжает работать."""
    from bot.handlers import dashboard

    async def broken(*_args: object, **_kwargs: object) -> bytes:
        raise RuntimeError("диск переполнен")

    h = app
    await seed_team(h)
    monkeypatch.setattr(dashboard.export_service, "build_report_xlsx", broken)
    await h.send_command(MGR, "export")
    log = await h.press_button(MGR, "Этот месяц")
    assert log.answers and not log.documents
    assert h.last_text(MGR) == "⚠️ Не удалось подготовить отчёт. Попробуйте ещё раз чуть позже."
    assert not any(isinstance(m, SendDocument) for m in log)

    await h.send_command(MGR, "team")
    assert lines(h.last_text(MGR))[0] == f"📊 Команда · {WEEK_LABEL}"


async def test_export_summary_covers_every_task_of_blocked_employee(app: BotHarness, clock: Clock) -> None:
    """Сидоров выполнил задачу на этой неделе и уволился (начальник его заблокировал).
    В выгрузке за неделю его задача есть на листе «Задачи», значит и на «Сводке» он должен быть,
    а «Итого по команде» — учитывать все задачи листа «Задачи» (иначе итог не сходится с таблицей)."""
    h = app
    team = await seed_team(h)
    await seed_tz_table(h, clock)
    await seed_done(h, clock, SIDOROV, "Бюджет на квартал", 50, 80, local(6))
    async with h.db() as s:
        await users.block_user(s, team.sidorov.id, await _user(s, MGR))
        await s.commit()

    await h.send_command(MGR, "export")
    [doc] = (await h.press_button(MGR, "Эта неделя")).documents
    task_rows = _sheet_rows(doc.content, "Задачи")[1:]
    assert any(row[1] == "Сидоров Пётр Ильич" for row in task_rows)
    summary = {row[0]: row for row in _sheet_rows(doc.content, "Сводка")[1:] if row and row[0]}
    assert "Сидоров Пётр Ильич" in summary
    assert summary["Итого по команде"][3] == len(task_rows)
    assert summary["Итого по команде"][2] == 94.3


async def test_blocked_employee_history_stays_reachable(app: BotHarness, clock: Clock) -> None:
    """Сидоров уволился, начальник его заблокировал. На дашборде команды его больше нет, но
    в «👥 Сотрудники» у его карточки есть «📊 Карточка»: эффективность за неделю (80 %) и
    история оценок («Бюджет на квартал») по-прежнему доступны — разблокировать для этого не нужно."""
    h = app
    team = await seed_team(h)
    await seed_done(h, clock, SIDOROV, "Бюджет на квартал", 50, 80, local(6))
    async with h.db() as s:
        await users.block_user(s, team.sidorov.id, await _user(s, MGR))
        await s.commit()

    await h.send_command(MGR, "team")
    assert "Сидоров" not in h.last_text(MGR)

    await h.send_command(MGR, "staff")
    await h.press_button(MGR, "Сидоров")
    assert h.buttons(MGR)[:2] == ["📊 Карточка", "🔓 Разблокировать"]
    log = await h.press_button(MGR, "📊 Карточка")
    assert log.answers and log.alert is None
    assert lines(h.last_text(MGR))[0] == "👤 Сидоров Пётр Ильич — 80 %"
    await h.press_button(MGR, "История оценок")
    assert "Бюджет на квартал" in h.last_text(MGR)
    assert (await h.get_user(SIDOROV)).status == models.UserStatus.BLOCKED


async def test_manager_typing_my_kpi_gets_pointed_to_team(app: BotHarness, clock: Clock) -> None:
    """Начальник (например, недавно повышенный из сотрудников) набирает /kpi или старую кнопку
    «📈 Моя эффективность»: вместо «Не понял» бот объясняет, что это раздел сотрудника, и
    показывает, где смотреть эффективность команды."""
    h = app
    await seed_team(h)
    await h.send_command(MGR, "start")
    for send in (h.send_command(MGR, "kpi"), h.send_text(MGR, BTN_MY_KPI)):
        log = await send
        assert "раздел сотрудника" in log.text and BTN_TEAM in log.text
        assert "Не понял" not in log.text


# --- Разное ---------------------------------------------------------------------------------------


async def test_team_button_interrupts_review_dialog(app: BotHarness, clock: Clock) -> None:
    """Начальник начал менять оценку («✏️ Изменить оценку»), но передумал и нажал «📊 Команда».
    Дашборд открывается, диалог проверки сброшен: следующее «90» оценкой не считается."""
    h = app
    await seed_team(h)
    task_id, _ = await seed_submitted(h, clock, IVANOV, "Анализ договоров", 20, local(6), ai=110)
    await h.send_command(MGR, "start")
    await h.press_menu(MGR, BTN_REVIEW)
    await h.press_button(MGR, "Анализ договоров")
    await h.press_button(MGR, "Изменить оценку")
    assert await h.get_state(MGR) is not None

    await h.press_menu(MGR, BTN_TEAM)
    assert lines(h.last_text(MGR))[0] == f"📊 Команда · {WEEK_LABEL}"
    assert await h.get_state(MGR) is None

    await h.send_text(MGR, "90")
    task = await h.get_task(task_id)
    assert task.status == TaskStatus.SUBMITTED and task.final_score is None


async def test_my_kpi_button_interrupts_proposal_dialog(app: BotHarness, clock: Clock) -> None:
    """Иванов начал вносить поручение, но нажал «📈 Моя эффективность»: открывается карточка,
    диалог сброшен — следующий текст не становится названием поручения."""
    h = app
    await seed_team(h)
    await h.send_command(IVANOV, "start")
    await h.press_menu(IVANOV, BTN_PROPOSE)
    assert await h.get_state(IVANOV) is not None

    await h.press_menu(IVANOV, BTN_MY_KPI)
    rows = lines(h.last_text(IVANOV))
    assert rows[0] == "👤 Иванов Иван Иванович — нет данных"
    assert "Задач в этом периоде нет." in rows
    assert await h.get_state(IVANOV) is None

    await h.send_text(IVANOV, "Подготовить справку по закупкам")
    assert await h.scalar(select(Task).where(Task.title.contains("справку"))) is None


async def test_html_special_characters_in_names_and_titles(app: BotHarness, clock: Clock) -> None:
    """ФИО, должность и название задачи с «<», «>» и «&» не ломают ни дашборд, ни карточку, ни историю
    (Telegram отклонил бы неэкранированный HTML — FakeSession проверяет это так же строго)."""
    h = app
    await h.seed_user(MGR, "Петрова Анна Сергеевна", role="manager")
    await h.seed_user(IVANOV, "Д'Артаньян Шарль <Гасконец>", position="Юрист & аналитик <главный>")
    await seed_done(h, clock, IVANOV, "Отчёт <черновик> & итог", 20, 100, local(6))

    await h.send_command(MGR, "team")
    assert "1. Д'Артаньян Ш. <. — 100 %" in lines(h.last_text(MGR))
    await h.press_button(MGR, "Д'Артаньян")
    rows = lines(h.last_text(MGR))
    assert rows[:2] == ["👤 Д'Артаньян Шарль <Гасконец> — 100 %", "💼 Юрист & аналитик <главный>"]
    assert "• Отчёт <черновик> & итог — вес 20 % × 100 %" in rows
    await h.press_button(MGR, "📜 История оценок")
    assert any(row.endswith("Отчёт <черновик> & итог") for row in lines(h.last_text(MGR)))

    await h.send_command(IVANOV, "kpi")
    assert lines(h.last_text(IVANOV))[0] == "👤 Д'Артаньян Шарль <Гасконец> — 100 %"


async def test_weekly_digest_buttons_lead_to_dashboard(app: BotHarness, clock: Clock) -> None:
    """Начальник получил от планировщика еженедельную сводку за прошлую неделю. Её кнопки работают
    как дашборд: «Позже ▶» — текущая неделя, «👤 Иванов» — карточка сотрудника."""
    from bot.scheduler import jobs

    h = app
    await seed_team(h)
    await seed_done(h, clock, IVANOV, "Прошлая неделя", 20, 80, local(1))
    await seed_done(h, clock, IVANOV, "Подготовка ТЗ", 30, 100, local(5))

    log = await h.capture(jobs.weekly_digest(h.bot, h.sessionmaker))
    assert PREV_WEEK_LABEL in log.to(MGR).text
    assert "Позже ▶" in h.buttons(MGR)

    await h.press_button(MGR, "Позже ▶")
    assert lines(h.last_text(MGR))[0] == f"📊 Команда · {WEEK_LABEL}"
    assert "1. Иванов И. И. — 100 %" in lines(h.last_text(MGR))
    await h.press_button(MGR, "👤 Иванов")
    assert lines(h.last_text(MGR))[0] == "👤 Иванов Иван Иванович — 100 %"
