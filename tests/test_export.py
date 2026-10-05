"""bot.services.export: отчёт Excel за период (SPEC 3.6)."""

from __future__ import annotations

from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from bot.db.models import TaskStatus, User
from bot.services import tasks as svc
from bot.services.export import build_report_xlsx
from bot.services.periods import get_period

SUMMARY_HEADERS = [
    "Сотрудник", "Должность", "KPI %", "Задач", "Выполнено", "В срок %", "Просрочено", "На проверке",
    "В работе", "Перевыполнено", "Внесено самостоятельно",
]
TASK_HEADERS = [
    "№", "Сотрудник", "Задача", "Ожидаемый результат", "План", "Факт", "Срок", "Сдано", "Просрочка дн.",
    "Вес %", "Приоритет", "Статус", "Оценка AI", "Итоговая оценка", "Решение", "Комментарий",
]
JOURNAL_HEADERS = ["Дата", "Задача", "Кто", "Событие", "Детали"]


def header(ws) -> list:
    return [cell.value for cell in ws[1]][: ws.max_column]


def row_by_first_cell(ws, value: str) -> list:
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == value:
            return list(row)
    raise AssertionError(f"нет строки «{value}»")


async def test_report_workbook(session, clock, manager: User, employee: User, employee2: User, make_task) -> None:
    # Задачи employee через сервисы: ТЗ-пример весов и оценок -> KPI 101.5.
    for weight, score in ((30, 100), (20, 110), (20, 90), (30, 105)):
        task = await svc.create_task(
            session, creator=manager, assignee_id=employee.id, title=f"=SUM(A1) задача {weight}",
            expected_result="Проверить 100 договоров", deadline=clock.now + timedelta(hours=4), weight=weight,
            plan_value=100, plan_unit="договоров",
        )
        sub = await svc.submit_result(session, task.id, employee, fact_text="Сделано", fact_value=score)
        await svc.record_evaluation(session, sub.id, score=score, rationale="r", source="rules")
        await svc.review_confirm(session, sub.id, manager)
    # employee2: одна просроченная несданная задача -> KPI 0.
    await make_task(employee2, deadline=clock.now - timedelta(hours=2), weight=10, title="Просроченная")

    period = get_period("week", 0, clock.now)
    data = await build_report_xlsx(session, period, clock.now)
    assert isinstance(data, bytes) and data[:2] == b"PK"

    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Сводка", "Задачи", "Журнал"]
    summary, tasks_ws, journal = wb["Сводка"], wb["Задачи"], wb["Журнал"]
    assert header(summary) == SUMMARY_HEADERS
    assert header(tasks_ws) == TASK_HEADERS
    assert header(journal) == JOURNAL_HEADERS

    for ws in (summary, tasks_ws, journal):
        assert all(cell.font.bold for cell in ws[1])
        assert ws.freeze_panes == "A2"

    first = row_by_first_cell(summary, employee.full_name)
    assert isinstance(first[2], (int, float)) and first[2] == pytest.approx(101.5)  # KPI — число, не «102 %»
    assert first[1] == "Юрист"
    assert first[3:6] == [4, 4, 100]
    second = row_by_first_cell(summary, employee2.full_name)
    assert second[2] == 0 and second[6] == 1  # просрочено

    task_rows = list(tasks_ws.iter_rows(min_row=2, values_only=True))
    assert len(task_rows) == 5
    titles = [row[2] for row in task_rows]
    assert "=SUM(A1) задача 30" in titles  # пользовательский текст не стал формулой
    overdue_row = next(row for row in task_rows if row[2] == "Просроченная")
    assert overdue_row[11] == "Просрочена"
    done_row = next(row for row in task_rows if row[2] == "=SUM(A1) задача 20")
    assert done_row[13] == 110 and done_row[4] == "100 договоров"

    journal_rows = list(journal.iter_rows(min_row=2, values_only=True))
    events = [row[3] for row in journal_rows]
    assert events.count("Задача поставлена") == 4
    assert "Сдан результат" in events and "Оценка подтверждена" in events


async def test_report_excludes_tasks_outside_period(session, clock, employee: User, make_task) -> None:
    await make_task(employee, deadline=clock.now - timedelta(days=30), status=TaskStatus.DONE, final_score=100)
    await make_task(employee, deadline=clock.now, status=TaskStatus.CANCELLED, title="Отменённая")
    period = get_period("week", 0, clock.now)
    wb = load_workbook(BytesIO(await build_report_xlsx(session, period, clock.now)))
    assert wb["Задачи"].max_row == 1
    row = row_by_first_cell(wb["Сводка"], employee.full_name)
    assert row[2] is None and row[3] == 0


async def test_report_on_empty_database(session) -> None:
    data = await build_report_xlsx(session, get_period("month", -1))
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Сводка", "Задачи", "Журнал"]
    assert header(wb["Задачи"]) == TASK_HEADERS
