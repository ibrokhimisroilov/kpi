"""Выгрузка отчёта за период в Excel: «Сводка», «Задачи», «Журнал».

Числа пишутся числами (KPI 101.5, а не «102 %»), даты — строками местного времени
«dd.mm.yyyy HH:MM». Файл собирается в памяти и возвращается байтами.

Язык отчёта — язык того, кто его запросил (bot.i18n, SPEC.md §14): названия листов и колонок, статусы,
события и подписи журнала проходят через ``_t``; тексты задач и ФИО остаются как написаны.
"""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from bot import i18n
from bot.config import get_settings
from bot.db.models import (
    EventType,
    Priority,
    ReviewDecision,
    Submission,
    Role,
    Task,
    TaskEvent,
    TaskStatus,
    User,
    UserStatus,
)
from bot.services.kpi import KpiResult, TaskSnapshot, compute_kpi, kpi_for_team, period_tasks, team_kpi
from bot.services.periods import Period
from bot.utils.dates import days_between, fmt_date, fmt_datetime, utcnow

STATUS_NAMES = {
    TaskStatus.PROPOSED: "На подтверждении",
    TaskStatus.ACTIVE: "В работе",
    TaskStatus.SUBMITTED: "На проверке",
    TaskStatus.REWORK: "На доработке",
    TaskStatus.DONE: "Выполнена",
    TaskStatus.CANCELLED: "Отменена",
    TaskStatus.REJECTED: "Отклонена",
}
PRIORITY_NAMES = {Priority.HIGH: "Высокий", Priority.MEDIUM: "Средний", Priority.LOW: "Низкий"}
DECISION_NAMES = {
    ReviewDecision.APPROVED: "Оценка AI подтверждена",
    ReviewDecision.CHANGED: "Оценка изменена",
    ReviewDecision.REWORK: "Возвращено на доработку",
}
EVENT_NAMES = {
    EventType.CREATED: "Задача поставлена",
    EventType.PROPOSED: "Поручение внесено сотрудником",
    EventType.APPROVED: "Поручение подтверждено",
    EventType.REJECTED: "Поручение отклонено",
    EventType.ACCEPTED: "Принята в работу",
    EventType.EDITED: "Задача изменена",
    EventType.SUBMITTED: "Сдан результат",
    EventType.AI_EVALUATED: "Предварительная оценка",
    EventType.SCORE_CONFIRMED: "Оценка подтверждена",
    EventType.SCORE_CHANGED: "Оценка изменена",
    EventType.REWORK: "Возвращена на доработку",
    EventType.CANCELLED: "Задача отменена",
    EventType.REMINDER: "Отправлено напоминание",
    EventType.WEIGHT_SUGGESTED: "Предложен вес поручения",
}
AUTO_DECISION_NAME = "Оценка AI подтверждена автоматически"  # начальник не ответил вовремя (bot.services.auto)
# Подписи ключей TaskEvent.data для колонки «Детали».
_DATA_LABELS = {
    "title": "Название",
    "expected_result": "Ожидаемый результат",
    "description": "Описание",
    "plan_value": "План",
    "plan_unit": "Единица",
    "deadline": "Срок",
    "new_deadline": "Новый срок",
    "old_deadline": "Прежний срок",
    "previous_status": "Прежний статус",
    "files": "Файлов",
    "priority": "Приоритет",
    "weight": "Вес",
    "reason": "Причина",
    "comment": "Комментарий",
    "score": "Оценка",
    "ai_score": "Оценка AI",
    "final_score": "Итоговая оценка",
    "attempt": "Попытка",
    "source": "Источник",
    "model": "Модель",
    "rationale": "Обоснование",
    "late_days": "Просрочка, дн.",
    "is_late": "С опозданием",
    "kind": "Напоминание",
    "auto": "Автоматически",
    "previous": "Прежняя оценка",
    "after_auto": "После автоподтверждения",
    "note": "Пояснение",
}
# Служебные ключи (id записей) в отчёт не выводятся.
_HIDDEN_KEYS = frozenset({"submission_id", "assignee_id", "task_id"})
_PERCENT_KEYS = frozenset({"score", "ai_score", "final_score", "weight", "previous"})
_SOURCE_NAMES = {"ai": "AI", "rules": "правила", "manager": "начальник", "employee": "сотрудник"}
_DETAILS_LIMIT = 500

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="DDEBF7")
_HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)
_TOP = Alignment(vertical="top")
_TOP_WRAP = Alignment(vertical="top", wrap_text=True)
_MIN_WIDTH, _MAX_WIDTH, _MAX_WRAP_WIDTH = 6, 40, 60


def _t(text: str | None) -> str | None:
    """Подпись отчёта (не слова пользователя) — на язык запросившего."""
    return i18n.tr(text) if text else text


@dataclass(frozen=True)
class _Col:
    title: str
    number_format: str | None = None  # для числовых колонок
    wrap: bool = False                 # длинный текст — перенос по словам


_SUMMARY_COLUMNS = (
    _Col("Сотрудник"),
    _Col("Должность", wrap=True),
    _Col("KPI %", "0.0"),
    _Col("Задач"),
    _Col("Выполнено"),
    _Col("В срок %", "0.0"),
    _Col("Просрочено"),
    _Col("На проверке"),
    _Col("В работе"),
    _Col("Перевыполнено"),
    _Col("Внесено самостоятельно"),
)
_TASK_COLUMNS = (
    _Col("№"),
    _Col("Сотрудник"),
    _Col("Задача", wrap=True),
    _Col("Ожидаемый результат", wrap=True),
    _Col("План", wrap=True),
    _Col("Факт", wrap=True),
    _Col("Срок"),
    _Col("Сдано"),
    _Col("Просрочка дн.", "0.0"),
    _Col("Вес %"),
    _Col("Приоритет"),
    _Col("Статус"),
    _Col("Оценка AI"),
    _Col("Итоговая оценка"),
    _Col("Решение"),
    _Col("Комментарий", wrap=True),
)
_JOURNAL_COLUMNS = (
    _Col("Дата"),
    _Col("Задача", wrap=True),
    _Col("Кто"),
    _Col("Событие"),
    _Col("Детали", wrap=True),
)


async def build_report_xlsx(session: AsyncSession, period: Period, now: datetime | None = None) -> bytes:
    """Отчёт за период: сводка по сотрудникам, задачи периода и журнал их событий."""
    now = now or utcnow()
    team = await kpi_for_team(session, period, now)
    tasks = sorted(
        await period_tasks(session, period),
        key=lambda t: (t.assignee.full_name, t.deadline, t.id),
    )
    team = team + _former_members(team, tasks, now)
    events = await _events(session, [t.id for t in tasks])

    wb = Workbook()
    wb.properties.title = f"{_t('Эффективность')} — {_t(period.label)}"
    summary = wb.active
    summary.title = _t("Сводка")
    _fill_summary(summary, period, team, now)
    _write_table(wb.create_sheet(_t("Задачи")), _TASK_COLUMNS, [_task_row(t, now) for t in tasks])
    titles = {t.id: f"#{t.id} {t.title}" for t in tasks}
    _write_table(wb.create_sheet(_t("Журнал")), _JOURNAL_COLUMNS, [_event_row(e, titles) for e in events])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _events(session: AsyncSession, task_ids: list[int]) -> list[TaskEvent]:
    if not task_ids:
        return []
    stmt = (
        select(TaskEvent)
        .where(TaskEvent.task_id.in_(task_ids))
        .order_by(TaskEvent.created_at, TaskEvent.id)
    )
    return list((await session.scalars(stmt)).all())


# --- Сводка ---------------------------------------------------------------------------


def _former_members(
    team: list[tuple[User, KpiResult]], tasks: Sequence[Task], now: datetime
) -> list[tuple[User, KpiResult]]:
    """Исполнители задач периода, которых нет среди активных сотрудников (заблокирован, стал начальником).

    Их задачи есть на листе «Задачи», поэтому они нужны и в «Сводке» — иначе «Итого по команде»
    не сходится с таблицей задач. Дашборд бота показывает только действующих сотрудников.
    """
    known = {user.id for user, _ in team}
    by_assignee: dict[int, list[Task]] = defaultdict(list)
    for task in tasks:
        if task.assignee_id not in known:
            by_assignee[task.assignee_id].append(task)
    overdue_as_zero = get_settings().overdue_counts_as_zero
    rows = [
        (items[0].assignee, compute_kpi([TaskSnapshot.from_task(t) for t in items], now, overdue_as_zero))
        for items in by_assignee.values()
    ]
    return sorted(rows, key=lambda row: (row[0].full_name, row[0].id))


def _position(user: User) -> str | None:
    """Должность; для бывших участников команды — пометка, почему их нет на дашборде."""
    if user.status == UserStatus.BLOCKED:
        note = "заблокирован"
    elif user.status != UserStatus.ACTIVE:
        note = "не подтверждён"
    elif user.role == Role.MANAGER:
        note = "сейчас начальник"
    else:
        return user.position
    note = _t(note)
    return f"{user.position} ({note})" if user.position else f"({note})"


def _fill_summary(ws: Worksheet, period: Period, team: list[tuple[User, KpiResult]], now: datetime) -> None:
    rows = [_summary_row(user, res) for user, res in team]
    _write_table(ws, _SUMMARY_COLUMNS, rows, total=_summary_total(team))
    last_day = period.end - timedelta(seconds=1)
    ws.append([])
    ws.append([_t(f"Период: {period.label} ({fmt_date(period.start)} – {fmt_date(last_day)})")])
    ws.append([_t(f"Сформировано: {fmt_datetime(now)}")])
    note = "KPI % = Σ(вес × оценка) / Σ(вес) по оценённым задачам периода"
    if get_settings().overdue_counts_as_zero:
        note += "; просроченные несданные задачи учитываются как 0 %"
    ws.append([_t(note + ".")])


def _summary_row(user: User, res: KpiResult) -> list[Any]:
    return [
        user.full_name,
        _position(user),
        _round(res.kpi),
        res.total,
        res.done,
        _round(res.on_time_pct),
        res.overdue_total,
        res.on_review,
        res.in_progress,
        res.overperformed,
        res.self_initiated,
    ]


def _summary_total(team: list[tuple[User, KpiResult]]) -> list[Any]:
    results = [res for _, res in team]
    done = sum(r.done for r in results)
    on_time = sum(r.done_on_time for r in results)
    return [
        _t("Итого по команде"),
        None,
        _round(team_kpi(team)),
        sum(r.total for r in results),
        done,
        _round(on_time / done * 100) if done else None,
        sum(r.overdue_total for r in results),
        sum(r.on_review for r in results),
        sum(r.in_progress for r in results),
        sum(r.overperformed for r in results),
        sum(r.self_initiated for r in results),
    ]


# --- Задачи ---------------------------------------------------------------------------


def _task_row(task: Task, now: datetime) -> list[Any]:
    last = task.last_submission
    return [
        task.id,
        task.assignee.full_name,
        task.title,
        task.expected_result,
        _value_with_unit(task.plan_value, task.plan_unit),
        _fact_text(last, task.plan_unit),
        fmt_datetime(task.deadline),
        fmt_datetime(last.created_at) if last else None,
        _late_days(task, last, now),
        task.weight,
        _t(PRIORITY_NAMES.get(task.priority, str(task.priority))),
        _t(_status_name(task, now)),
        task.ai_score,
        task.final_score,
        _t(_decision_name(last)),
        last.review_comment if last else None,
    ]


def _status_name(task: Task, now: datetime) -> str:
    if task.is_open and task.deadline < now:
        return "Просрочена"
    return STATUS_NAMES.get(task.status, str(task.status))


def _late_days(task: Task, last: Submission | None, now: datetime) -> float:
    """Просрочка в днях: для открытых — на текущий момент, для сданных — на момент сдачи."""
    if task.is_open:
        return round(max(0.0, days_between(now, task.deadline)), 1)
    return round(last.late_days, 1) if last else 0.0


def _fact_text(last: Submission | None, unit: str | None) -> str | None:
    """Последняя сдача: значение, что сделано, результат, файлы."""
    if last is None:
        return None
    lines = [_value_with_unit(last.fact_value, unit), last.fact_text]
    if last.result_text:
        lines.append(f"{_t('Результат')}: {last.result_text}")
    if last.attachments:
        names = ", ".join(a.file_name or str(a.kind) for a in last.attachments)
        lines.append(f"{_t('Файлы')}: {names}")
    return "\n".join(line for line in lines if line)


def _value_with_unit(value: float | None, unit: str | None) -> str | None:
    if value is None:
        return None
    return f"{_num(value)} {unit}" if unit else _num(value)


# --- Журнал ---------------------------------------------------------------------------


def _event_row(event: TaskEvent, titles: dict[int, str]) -> list[Any]:
    return [
        fmt_datetime(event.created_at),
        titles.get(event.task_id, f"#{event.task_id}"),
        event.actor.full_name if event.actor else _t("Система"),
        _t(EVENT_NAMES.get(event.type, str(event.type))),
        _event_details(event.data or {}),
    ]


def _event_details(data: dict[str, Any]) -> str | None:
    """Данные события по-русски: «Срок: 05.10.2026 18:00 → 07.10.2026 18:00»."""
    lines: list[str] = []
    for key, value in data.items():
        if key in _HIDDEN_KEYS or value is None or value == "":
            continue
        if key == "changes" and isinstance(value, dict):
            lines += [_change_line(field, change) for field, change in value.items()]
        else:
            lines.append(f"{_t(_DATA_LABELS.get(key, key))}: {_data_value(key, value)}")
    text = "\n".join(lines)
    return _clip(text) if text else None


def _change_line(field: str, change: Any) -> str:
    label = _t(_DATA_LABELS.get(field, field))
    if isinstance(change, (list, tuple)) and len(change) == 2:
        old, new = change
        return f"{label}: {_data_value(field, old)} → {_data_value(field, new)}"
    return f"{label}: {_data_value(field, change)}"


def _decision_name(sub: Submission | None) -> str | None:
    """Решение по последней сдаче для отчёта (None — решения ещё нет)."""
    if sub is None or not sub.decision:
        return None
    if sub.auto_confirmed:
        return AUTO_DECISION_NAME
    return DECISION_NAMES.get(sub.decision)


def _data_value(key: str, value: Any) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, bool):
        return _t("да" if value else "нет") or ""
    if isinstance(value, (int, float)):
        return _num(value) + (" %" if key in _PERCENT_KEYS else "")
    text = str(value)
    if key == "priority":
        return _t(PRIORITY_NAMES.get(text)) or text
    if key == "previous_status":
        return _t(STATUS_NAMES.get(text)) or text
    if key == "source":
        return _t(_SOURCE_NAMES.get(text)) or text
    if key == "kind":
        return _t(_reminder_name(text)) or text
    if "deadline" in key:
        return _iso_to_local(text)
    return _clip(text)


def _iso_to_local(text: str) -> str:
    """ISO-дата UTC из журнала -> «05.10.2026 18:00» местного времени."""
    try:
        return fmt_datetime(datetime.fromisoformat(text))
    except ValueError:
        return text


def _reminder_name(kind: str) -> str:
    """Ключ ReminderLog -> понятное название напоминания."""
    fixed = {
        "before_hours": "в день срока",
        "deadline_passed": "срок истёк",
        "overdue_manager": "начальнику о просрочке",
    }
    if kind in fixed:
        return fixed[kind]
    days = kind.removeprefix("before_").removesuffix("d")
    if kind.startswith("before_") and days.isdigit():
        return f"за {days} дн. до срока"
    if kind.startswith("overdue_"):
        return "ежедневное о просрочке"
    if kind.startswith("review_"):
        return "о непроверенном результате"
    return kind


# --- Общее оформление -------------------------------------------------------------------


def _write_table(
    ws: Worksheet,
    columns: Sequence[_Col],
    rows: Sequence[Sequence[Any]],
    *,
    total: Sequence[Any] | None = None,
) -> None:
    """Таблица с жирной шапкой, закреплённой первой строкой, фильтром и шириной по содержимому."""
    ws.append([_t(col.title) for col in columns])
    for row in rows:
        ws.append(list(row))
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    if total is not None:
        ws.append(list(total))
        for cell in ws[ws.max_row]:
            cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"

    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN

    table_rows = [*rows, *([total] if total is not None else [])]
    for index, col in enumerate(columns, start=1):
        letter = get_column_letter(index)
        for cell in ws[letter][1:]:
            if cell.data_type == "f":
                cell.data_type = "s"  # текст пользователя «=…» — не формула
            cell.alignment = _TOP_WRAP if col.wrap else _TOP
            if col.number_format and isinstance(cell.value, (int, float)):
                cell.number_format = col.number_format
        values = [row[index - 1] for row in table_rows]
        ws.column_dimensions[letter].width = _column_width(col, values)


def _column_width(col: _Col, values: Sequence[Any]) -> float:
    longest = max(
        (len(line) for value in values if value is not None for line in _display(value).splitlines()),
        default=0,
    )
    limit = _MAX_WRAP_WIDTH if col.wrap else _MAX_WIDTH
    return min(max(longest, len(col.title), _MIN_WIDTH) + 2, limit)


def _display(value: Any) -> str:
    return _num(value) if isinstance(value, float) else str(value)


def _round(value: float | None) -> float | None:
    return round(value, 1) if value is not None else None


def _num(value: float) -> str:
    """100.0 -> «100», 2.5 -> «2,5»."""
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".").replace(".", ",")


def _clip(text: str) -> str:
    return text if len(text) <= _DETAILS_LIMIT else text[: _DETAILS_LIMIT - 1] + "…"
