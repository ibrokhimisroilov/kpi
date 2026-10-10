"""Извлечение содержимого файлов-подтверждений для AI-оценки.

PDF и изображения передаются AI как байты (types.Part), Word/Excel/текст — как извлечённый
текст. Всё остальное (видео, архивы, старые .doc/.xls) помечается «skipped» с причиной,
чтобы AI знал, что файл был, но прочитать его не удалось.

Gemini читает PDF и изображения сам. Запасным провайдерам (bot.ai.openai_compat) изображения уходят
data-URL только моделям, которые их видят; PDF и изображения для остальных моделей заменяются
пометкой «файл приложен, содержимое не передано». Текст файла — base.DataText: провайдер с маленьким
лимитом может сократить его середину, блок «<<< … >>>» при этом остаётся закрытым.
"""

from __future__ import annotations

import asyncio
import io
import logging
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePath
from typing import Any

from aiogram import Bot
from docx import Document
from docx.table import Table
from google.genai import types
from openpyxl import load_workbook

from bot.ai.base import TRIM_FILES, DataText
from bot.config import get_settings
from bot.db.models import Attachment, AttachmentKind

__all__ = ["EvidenceItem", "collect_evidence", "evidence_to_parts", "defuse_markers"]

# Границы блока данных сотрудника в запросе к AI — строки «<<<» и «>>>». Такие же цепочки в тексте
# сотрудника заменяются похожими «‹‹‹» / «›››», чтобы он не мог «закрыть» блок и дописать «инструкцию».
_MARKER_RE = re.compile(r"<{3,}|>{3,}")


def defuse_markers(text: str) -> str:
    """Обезвредить «<<<» и «>>>» в данных сотрудника (текст, имена файлов, содержимое файлов)."""
    return _MARKER_RE.sub(lambda m: ("‹" if m.group()[0] == "<" else "›") * len(m.group()), text)

logger = logging.getLogger(__name__)

MAX_TEXT_PER_FILE = 15_000       # символов текста с одного файла
MAX_TEXT_TOTAL = 60_000          # символов текста со всех файлов
MAX_INLINE_BYTES = 14 * 1024 * 1024  # PDF/фото в одном запросе (лимит Gemini ~20 МБ после base64)
BOT_API_LIMIT_MB = 20            # Bot API не отдаёт файлы больше 20 МБ
MAX_XLSX_ROWS = 200              # строк на лист Excel
DOWNLOAD_TIMEOUT_SEC = 60
_MIN_TEXT_CHUNK = 200            # меньше этого остатка текстового бюджета — файл не передаём
_TRUNCATED_MARK = "\n…[текст обрезан]"
_BINARY_KINDS = ("pdf", "image")  # передаются Gemini как байты

_MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Формат файла: (вид обработки, MIME для Gemini).
_BY_EXTENSION: dict[str, tuple[str, str]] = {
    ".pdf": ("pdf", "application/pdf"),
    ".jpg": ("image", "image/jpeg"),
    ".jpeg": ("image", "image/jpeg"),
    ".png": ("image", "image/png"),
    ".webp": ("image", "image/webp"),
    ".docx": ("docx", _MIME_DOCX),
    ".xlsx": ("xlsx", _MIME_XLSX),
    ".xlsm": ("xlsx", _MIME_XLSX),
    ".txt": ("plain", "text/plain"),
    ".csv": ("plain", "text/csv"),
    ".md": ("plain", "text/markdown"),
    ".json": ("plain", "application/json"),
}
_BY_MIME: dict[str, tuple[str, str]] = {
    "application/pdf": ("pdf", "application/pdf"),
    "image/jpeg": ("image", "image/jpeg"),
    "image/jpg": ("image", "image/jpeg"),
    "image/png": ("image", "image/png"),
    "image/webp": ("image", "image/webp"),
    _MIME_DOCX: ("docx", _MIME_DOCX),
    _MIME_XLSX: ("xlsx", _MIME_XLSX),
    "text/plain": ("plain", "text/plain"),
    "text/csv": ("plain", "text/csv"),
    "text/markdown": ("plain", "text/markdown"),
    "application/json": ("plain", "application/json"),
}
# Подписи форматов для AI: «Файл 1: Analysis.xlsx (Excel)».
_LABEL_BY_EXTENSION = {
    ".pdf": "PDF", ".jpg": "изображение", ".jpeg": "изображение", ".png": "изображение",
    ".webp": "изображение", ".gif": "изображение", ".heic": "изображение", ".docx": "Word",
    ".doc": "Word", ".xlsx": "Excel", ".xlsm": "Excel", ".xls": "Excel", ".txt": "текст",
    ".csv": "CSV", ".md": "Markdown", ".json": "JSON", ".pptx": "PowerPoint", ".ppt": "PowerPoint",
    ".zip": "архив", ".rar": "архив", ".7z": "архив", ".mp4": "видео", ".mov": "видео",
}
_LABEL_BY_MIME_PREFIX = {"image/": "изображение", "video/": "видео", "audio/": "аудио", "text/": "текст"}
_LABEL_BY_KIND = {"pdf": "PDF", "image": "изображение", "text": "текст"}


@dataclass
class EvidenceItem:
    name: str
    kind: str                     # "text" | "pdf" | "image" | "skipped"
    text: str | None = None       # извлечённый текст (обрезан до 15 000 симв.)
    data: bytes | None = None     # для pdf/image
    mime_type: str | None = None
    note: str | None = None       # почему пропущен


async def collect_evidence(bot: Bot, attachments: Sequence[Attachment]) -> list[EvidenceItem]:
    """Скачать и разобрать вложения сдачи. Никогда не бросает: проблемный файл -> "skipped"."""
    try:
        return await _collect(bot, attachments)
    except Exception:  # noqa: BLE001 - оценка должна пройти и без файлов
        logger.exception("Не удалось подготовить файлы для AI")
        return [
            _skipped(_display_name(att, i), att.mime_type, "не удалось прочитать файл")
            for i, att in enumerate(attachments, 1)
        ]


def evidence_to_parts(items: list[EvidenceItem]) -> list:
    """Превратить вложения в части запроса AI: строки и types.Part.from_bytes(...) (bot.ai.provider)."""
    parts: list = []
    for number, item in enumerate(items, 1):
        label = f"Файл {number}: {defuse_markers(item.name)} ({_format_label(item)})"
        if item.kind == "text" and item.text:
            # Запасной провайдер с маленьким лимитом может сократить середину текста (base.DataText).
            text = f"{label}. Текст файла (данные от сотрудника):\n<<<\n{defuse_markers(item.text)}\n>>>"
            parts.append(DataText(text, TRIM_FILES))
        elif item.kind in _BINARY_KINDS and item.data:
            mime_type = item.mime_type or "application/octet-stream"
            parts.append(f"{label}. Содержимое файла — следующей частью.")
            parts.append(types.Part.from_bytes(data=item.data, mime_type=mime_type))
        else:
            parts.append(f"{label} — не передан на анализ: {item.note or 'формат не поддерживается'}.")
    return parts


async def _collect(bot: Bot, attachments: Sequence[Attachment]) -> list[EvidenceItem]:
    settings = get_settings()
    max_mb = min(settings.ai_max_file_mb, BOT_API_LIMIT_MB)
    text_budget = MAX_TEXT_TOTAL
    bytes_budget = MAX_INLINE_BYTES
    items: list[EvidenceItem] = []

    for number, att in enumerate(attachments, 1):
        name = _display_name(att, number)
        if not settings.ai_read_files:
            items.append(_skipped(name, att.mime_type, "чтение файлов AI отключено в настройках"))
            continue
        item = await _process_one(bot, att, name, max_mb, text_budget, bytes_budget)
        if item.kind == "text" and item.text:
            text_budget -= len(item.text)
        elif item.data:
            bytes_budget -= len(item.data)
        items.append(item)
    return items


async def _process_one(
    bot: Bot, att: Attachment, name: str, max_mb: int, text_budget: int, bytes_budget: int
) -> EvidenceItem:
    """Обработать одно вложение с учётом лимитов размера и общего бюджета."""
    fmt = _detect_format(att, name)
    if fmt is None:
        return _skipped(name, att.mime_type, _unsupported_reason(att, name))
    handler, mime = fmt
    limit_bytes = max_mb * 1024 * 1024
    if att.file_size and att.file_size > limit_bytes:
        return _skipped(name, mime, f"файл больше {max_mb} МБ")
    if handler in _BINARY_KINDS and att.file_size and att.file_size > bytes_budget:
        return _skipped(name, mime, "превышен общий объём файлов для анализа")
    if handler not in _BINARY_KINDS and text_budget < _MIN_TEXT_CHUNK:
        return _skipped(name, mime, "превышен общий объём текста для анализа")

    try:
        data = await _download(bot, att.file_id)
    except Exception as exc:  # noqa: BLE001 - сеть/Telegram: файл просто не попадёт в анализ
        logger.warning("Не удалось скачать файл %r: %s", name, type(exc).__name__)
        return _skipped(name, mime, "не удалось скачать файл из Telegram")
    if data is None:
        return _skipped(name, mime, "не удалось скачать файл из Telegram")
    if len(data) > limit_bytes:
        return _skipped(name, mime, f"файл больше {max_mb} МБ")

    if handler in _BINARY_KINDS:
        if len(data) > bytes_budget:
            return _skipped(name, mime, "превышен общий объём файлов для анализа")
        return EvidenceItem(name=name, kind=handler, data=data, mime_type=mime)

    try:
        text = await asyncio.to_thread(_EXTRACTORS[handler], data)
    except Exception as exc:  # noqa: BLE001 - повреждённый/зашифрованный файл
        logger.warning("Не удалось извлечь текст из %r: %s", name, type(exc).__name__)
        return _skipped(name, mime, "файл повреждён или защищён, текст извлечь не удалось")
    text = text.strip()
    if not text:
        return _skipped(name, mime, "в файле не найден текст")
    return EvidenceItem(name=name, kind="text", text=_limit_text(text, text_budget), mime_type=mime)


async def _download(bot: Bot, file_id: str) -> bytes | None:
    """Скачать файл по file_id в память (Bot.download без destination возвращает BytesIO)."""
    buffer = await bot.download(file_id, timeout=DOWNLOAD_TIMEOUT_SEC)
    if buffer is None:
        return None
    if isinstance(buffer, io.BytesIO):
        return buffer.getvalue()
    return buffer.read()


def _detect_format(att: Attachment, name: str) -> tuple[str, str] | None:
    """(вид обработки, MIME) по типу вложения, MIME и расширению; None — не поддерживается."""
    if att.kind == AttachmentKind.PHOTO:
        return "image", "image/jpeg"  # Telegram пережимает фото в JPEG
    if att.kind == AttachmentKind.VIDEO:
        return None
    mime = (att.mime_type or "").split(";")[0].strip().lower()
    if mime in _BY_MIME:
        return _BY_MIME[mime]
    return _BY_EXTENSION.get(PurePath(name).suffix.lower())


def _unsupported_reason(att: Attachment, name: str) -> str:
    suffix = PurePath(name).suffix.lower()
    mime = (att.mime_type or "").lower()
    if att.kind == AttachmentKind.VIDEO or mime.startswith("video/"):
        return "видео не анализируется, его посмотрит начальник"
    if suffix in (".doc", ".xls", ".ppt"):
        return "старый формат Office — сохраните файл как .docx/.xlsx"
    return "формат не поддерживается для анализа"


def _extract_docx(data: bytes) -> str:
    """Текст документа Word: абзацы и таблицы в порядке следования."""
    lines: list[str] = []
    for block in Document(io.BytesIO(data)).iter_inner_content():
        if isinstance(block, Table):
            lines.extend(_docx_table_rows(block))
        elif block.text.strip():
            lines.append(block.text.strip())
    return "\n".join(lines)


def _docx_table_rows(table: Table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            value = " ".join(cell.text.split())
            if not cells or cells[-1] != value:  # объединённые ячейки повторяются — схлопываем
                cells.append(value)
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def _extract_xlsx(data: bytes) -> str:
    """Значения ячеек Excel (кэшированные результаты формул), до 200 строк на лист."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sections = [_xlsx_sheet_text(sheet) for sheet in workbook.worksheets]
    finally:
        workbook.close()
    return "\n\n".join(section for section in sections if section)


def _xlsx_sheet_text(sheet: Any) -> str:
    """Лист Excel построчно: «A | B | C»; пустые строки и хвостовые пустые ячейки отброшены."""
    lines: list[str] = []
    truncated = False
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= MAX_XLSX_ROWS:
            truncated = True
            break
        values = [_cell_text(value) for value in row]
        while values and not values[-1]:
            values.pop()
        if values:
            lines.append(" | ".join(values))
    if not lines:
        return ""
    header = f"Лист «{sheet.title}»:"
    footer = [f"… (показаны первые {MAX_XLSX_ROWS} строк)"] if truncated else []
    return "\n".join([header, *lines, *footer])


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, float):
        return f"{value:.10g}"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y" if value.time() == time(0) else "%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return " ".join(str(value).split())


def _extract_plain(data: bytes) -> str:
    """Текстовый файл: UTF-8 (в т.ч. с BOM), иначе Windows-1251."""
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


_EXTRACTORS = {"docx": _extract_docx, "xlsx": _extract_xlsx, "plain": _extract_plain}


def _limit_text(text: str, budget: int) -> str:
    limit = min(MAX_TEXT_PER_FILE, budget)
    if len(text) <= limit:
        return text
    return text[: max(limit - len(_TRUNCATED_MARK), 0)] + _TRUNCATED_MARK


def _display_name(att: Attachment, number: int) -> str:
    if att.file_name:
        return att.file_name
    if att.kind == AttachmentKind.PHOTO:
        return f"Фото {number}"
    if att.kind == AttachmentKind.VIDEO:
        return f"Видео {number}"
    return f"Файл без имени {number}"


def _skipped(name: str, mime_type: str | None, note: str) -> EvidenceItem:
    return EvidenceItem(name=name, kind="skipped", mime_type=mime_type, note=note)


def _format_label(item: EvidenceItem) -> str:
    """Подпись формата для AI: Excel, Word, PDF, изображение, видео…"""
    suffix = PurePath(item.name).suffix.lower()
    if suffix in _LABEL_BY_EXTENSION:
        return _LABEL_BY_EXTENSION[suffix]
    mime = (item.mime_type or "").lower()
    if mime == "application/pdf":
        return "PDF"
    for prefix, label in _LABEL_BY_MIME_PREFIX.items():
        if mime.startswith(prefix):
            return label
    return _LABEL_BY_KIND.get(item.kind, "файл")
